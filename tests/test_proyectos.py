from pathlib import Path

import pytest
from openpyxl import load_workbook

from inmoscraper.exportar import guardar_csv
from inmoscraper.modelos import Anuncio
from inmoscraper.portales import detectar
from inmoscraper.portales.icasas import Icasas, url_icasas, url_icasas_pa
from inmoscraper.portales.mercadolibre import MercadoLibre
from inmoscraper.portales.supercasas import SuperCasas
from inmoscraper.proyectos import (Estado, cargar_catalogo, ruta_unidad, seleccionar, zonas_unicas)
from inmoscraper.ejecutor import Resultado
from inmoscraper.excel_proyectos import escribir_excels

FIX = Path(__file__).parent / "fixtures"
RAIZ = Path(__file__).parent.parent


def test_urls_icasas():
    assert url_icasas("venta", "casas", "jalisco", "tlajomulco-zuniga") == \
        "https://www.icasas.mx/venta/habitacionales-casas-jalisco-tlajomulco-zuniga-2_5_14_0_625_0"
    assert url_icasas_pa("venta", "apartamentos", "arraijan") == \
        "https://www.icasas.com.pa/venta/apartamentos/arraijan/list"


def test_icasas_panama_paginacion_y_moneda():
    p = detectar("https://www.icasas.com.pa/venta/apartamentos/arraijan")
    assert isinstance(p, Icasas) and p.moneda_defecto == "USD"
    assert p.url_pagina(1).endswith("/arraijan/list")  # p_1 da 404 en el sitio
    assert p.url_pagina(3).endswith("/arraijan/list/p_3")
    assert Icasas("https://www.icasas.com.pa/venta/casas/panama/list/p_7").url_pagina(2).endswith("/panama/list/p_2")


def test_supercasas_parseo_html_real():
    p = detectar("https://www.supercasas.com/buscar/?Locations=10005")
    assert isinstance(p, SuperCasas)
    assert p.url_pagina(1) == "https://www.supercasas.com/buscar/?Locations=10005"
    assert p.url_pagina(3).endswith("PagingPageSkip=2")
    a, b, c = p.parsear((FIX / "supercasas_listado.html").read_text(encoding="utf-8"), "https://www.supercasas.com/buscar/")
    assert a.id == "1446384" and a.url == "https://www.supercasas.com/apartamentos-venta-renacimiento/1446384/"
    assert (a.precio, a.moneda, a.operacion, a.tipo) == (360000.0, "USD", "venta", "Apartamento")
    assert (a.recamaras, a.banos, a.estacionamientos, a.superficie_m2) == (3, 3, 2, 170)


ML_HTML = """
<ol><li class="ui-search-layout__item"><div class="poly-card">
 <a class="poly-component__title" href="https://apartamento.mercadolibre.com.do/MRD-123456789-apartamento-en-cap-cana-_JM#polycard">Apartamento en Cap Cana 2 hab</a>
 <div class="poly-price__current"><span class="andes-money-amount"><span class="andes-money-amount__currency-symbol">US$</span>
   <span class="andes-money-amount__fraction">215.000</span></span></div>
 <ul class="poly-attributes-list"><li class="poly-attributes-list__item">85 m² útiles</li>
   <li class="poly-attributes-list__item">2 dormitorios</li><li class="poly-attributes-list__item">2 baños</li></ul>
 <span class="poly-component__location">Cap Cana, Punta Cana, La Altagracia</span>
</div></li></ol>
<span class="andes-pagination__page-count">de 42</span>
"""


def test_mercadolibre_paginacion_y_parseo():
    url = "https://inmuebles.mercadolibre.com.do/apartamentos/venta/cap-cana"
    p = detectar(url)
    assert isinstance(p, MercadoLibre) and p.pais == "do" and p.modo_preferido == "auto"
    assert p.url_pagina(1) == url + "/"
    assert p.url_pagina(2) == url + "/_Desde_49_NoIndex_True"
    assert MercadoLibre(p.url_pagina(3)).url_pagina(2).endswith("/cap-cana/_Desde_49_NoIndex_True")
    (a,) = p.parsear(ML_HTML, url)
    assert a.id == "MRD123456789" and "#" not in a.url
    assert (a.precio, a.moneda, a.superficie_m2, a.recamaras, a.banos) == (215000.0, "USD", 85, 2, 2)
    assert p.total_paginas(p.sopa(ML_HTML)) == 42


def test_bloqueo_por_redireccion_a_verificacion():
    from inmoscraper.descarga import parece_bloqueo
    assert parece_bloqueo(200, "<html></html>", "https://www.mercadolibre.com.pa/gz/account-verification?go=x")
    assert not parece_bloqueo(200, "<html>ok</html>", "https://inmuebles.mercadolibre.com.pa/casas/")


def test_catalogo_real_sin_repetidos():
    cat = cargar_catalogo(RAIZ / "proyectos.yaml")
    folios = [f for p in cat.proyectos for f in p.folios]
    assert len(folios) == len(set(folios)), "un folio no puede estar en dos proyectos"
    for p in cat.proyectos:
        for z in p.zonas:
            assert cat.unidades(z) is not None
    # zonas compartidas -> una sola descarga
    todos = list(cat.proyectos)
    ids = [u.id for z in zonas_unicas(todos) for u in cat.unidades(z)]
    assert len(ids) > len(set(ids))  # hay URLs repetidas entre zonas...
    comp = zonas_unicas(todos)
    assert {p.folios[0] for p in comp["sma"]} == {"4SR-BAJ-02631", "4SR-BAJ-02634"}
    assert [p.pais for p in seleccionar(cat, paises=["PA", "RD"])] == ["PA", "RD"]
    assert seleccionar(cat, solo=["4SR-PAN-02590"])[0].nombre == "COSTA PACIFICA"


def test_excel_replica_zona_compartida(tmp_path):
    yaml = tmp_path / "p.yaml"
    yaml.write_text("""
zonas:
  z1: {nombre: Zona Uno, pais: PA, fuentes: [{portal: icasas, lugar: arraijan, tipos: [casas]}]}
proyectos:
  - {folios: [A-1], nombre: Proyecto A, pais: PA, zonas: [z1]}
  - {folios: [B-2], nombre: Proyecto B, pais: PA, zonas: [z1]}
  - {folios: [C-3], nombre: Sin portales, pais: PA, zonas: []}
""", encoding="utf-8")
    cat = cargar_catalogo(yaml)
    (u,) = cat.unidades("z1")
    datos = tmp_path / "datos"
    guardar_csv([Anuncio(portal="icasas", url="https://x/1", id="1", titulo="Casa en X", precio=100000, moneda="USD",
                         superficie_m2=200, operacion="venta")], ruta_unidad(datos, u))
    Estado(datos).guardar(u, Resultado(zona=u.id, portal="icasas", url=u.url, anuncios=1, paginas=1, estado="ok"))
    archivos = escribir_excels(cat, cat.proyectos, datos, tmp_path / "out")
    assert len(archivos) == 4
    for nombre in ("A-1_proyecto-a.xlsx", "B-2_proyecto-b.xlsx"):
        wb = load_workbook(tmp_path / "out" / nombre)
        assert wb.sheetnames[:2] == ["Resumen", "Lista de precios"]
        fila = [c.value for c in wb["Lista de precios"][2]]
        assert 100000 in fila and 500.0 in fila  # precio y precio por m²
    wi = load_workbook(tmp_path / "out" / "00_indice_proyectos.xlsx")
    assert wi["Zonas únicas"].max_row == 2 and "Proyecto A" in wi["Zonas únicas"]["G2"].value

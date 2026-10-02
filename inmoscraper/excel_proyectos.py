"""Excel por proyecto (Resumen / Lista de precios / Hoteles / Ficha completa) + índice."""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .proyectos import (CacheFichas, Catalogo, Estado, Proyecto, Unidad, filas_unidad,
                        zonas_unicas)

_CAB = PatternFill("solid", fgColor="1F3864")
_FUENTE_CAB = Font(bold=True, color="FFFFFF")
_AVISO = PatternFill("solid", fgColor="FFF2CC")
_ERROR = PatternFill("solid", fgColor="F8CBAD")

PRECIOS = [  # (encabezado, clave, ancho)
    ("Zona", "zona_nombre", 26), ("Portal", "portal", 13), ("Operación", "operacion", 11),
    ("Tipo", "tipo", 16), ("Título", "titulo", 48), ("Precio", "precio", 14), ("Moneda", "moneda", 8),
    ("Sup. m²", "superficie_m2", 9), ("Sup. total m²", "superficie_total_m2", 11),
    ("Precio por m²", "precio_m2", 13), ("Recámaras", "recamaras", 10), ("Baños", "banos", 8),
    ("Estac.", "estacionamientos", 8), ("Ubicación", "ubicacion", 34), ("Anunciante", "anunciante", 26),
    ("Alerta", "alerta", 30), ("Link", "url", 50), ("Fecha scrape", "fecha_scrape", 20),
]
HOTELES = [
    ("Zona", "zona_nombre", 26), ("Hotel / alojamiento", "titulo", 46), ("Tipo", "tipo", 20),
    ("Precio (estancia)", "precio", 14), ("Moneda", "moneda", 8), ("Calificación", "calificacion", 11),
    ("Estrellas", "estrellas", 9), ("Opiniones", "opiniones", 28), ("Dirección", "direccion", 30),
    ("Distancia", "ubicacion", 22), ("Link", "url", 50), ("Fecha scrape", "fecha_scrape", 20),
]
NUMERICAS = {"precio", "superficie_m2", "superficie_total_m2", "precio_m2", "recamaras", "banos",
             "estacionamientos", "calificacion", "estrellas", "lat", "lon", "expensas"}
FICHA_BASE = [
    ("Zona", "zona_nombre", 26), ("Portal", "portal", 13), ("ID", "id", 18), ("Operación", "operacion", 11),
    ("Tipo", "tipo", 16), ("Título", "titulo", 44), ("Precio", "precio", 14), ("Moneda", "moneda", 8),
    ("Precio por m²", "precio_m2", 13), ("Sup. m²", "superficie_m2", 9), ("Sup. total m²", "superficie_total_m2", 11),
    ("Recámaras", "recamaras", 10), ("Baños", "banos", 8), ("Estac.", "estacionamientos", 8),
    ("Alerta", "alerta", 30),
    ("Ubicación", "ubicacion", 32), ("Dirección", "direccion", 28), ("Lat", "lat", 10), ("Lon", "lon", 10),
    ("Anunciante", "anunciante", 24), ("Descripción (listado)", "descripcion", 40), ("Link", "url", 50),
]


def _num(v: Any) -> Optional[float]:
    try:
        return float(v) if v not in ("", None) else None
    except (TypeError, ValueError):
        return None


def _alerta(fila: dict, precio, sup, pm2) -> str:
    """Marca (sin borrar) datos probablemente erróneos del anunciante, para filtrarlos en Excel."""
    if not precio:
        return "sin precio"
    venta = fila.get("operacion") == "venta"
    avisos = []
    if venta and precio < 5000 and fila.get("moneda") in ("USD", "EUR", ""):
        avisos.append("precio de venta atípico (muy bajo)")
    if venta and precio > 50_000_000:
        avisos.append("precio de venta atípico (muy alto)")
    if venta and pm2 and (pm2 < 30 or pm2 > 60_000):
        avisos.append("precio por m² atípico")
    return "; ".join(avisos)


def _enriquecer(fila: dict, u: Unidad, cat: Catalogo, fichas: CacheFichas) -> dict:
    fila = dict(fila)
    fila["zona_nombre"] = cat.zonas[u.zona].get("nombre", u.zona)
    fila["zona_clave"] = u.zona
    # precio por m²: sobre superficie construida; si el portal solo da la total (terrenos, EasyBroker) sobre esa
    p = _num(fila.get("precio"))
    s = _num(fila.get("superficie_m2")) or _num(fila.get("superficie_total_m2"))
    fila["precio_m2"] = round(p / s, 2) if p and s and s > 0 else None
    fila["alerta"] = _alerta(fila, p, s, fila["precio_m2"])
    try:
        extra = json.loads(fila.get("extra") or "{}")
    except json.JSONDecodeError:
        extra = {}
    fila["estrellas"] = extra.get("estrellas")
    fila["opiniones"] = extra.get("opiniones")
    # Algunos portales (InfoCasas) traen la ficha completa en el propio listado
    fila["_ficha"] = fichas.datos.get(fila.get("url", "")) or (extra.get("_ficha") if isinstance(extra, dict) else None) or {}
    return fila


def _hoja(wb: Workbook, nombre: str, columnas: list[tuple[str, str, int]], filas: list[dict]):
    ws = wb.create_sheet(nombre)
    ws.append([c[0] for c in columnas])
    for c in ws[1]:
        c.fill, c.font = _CAB, _FUENTE_CAB
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for f in filas:
        fila = []
        for _, k, _ in columnas:
            v = f.get(k)
            if k in NUMERICAS:
                v = _num(v)
            elif v in ("", None):
                v = None
            fila.append(v)
        ws.append(fila)
    for i, (_, k, ancho) in enumerate(columnas, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
        if k in ("precio", "expensas"):
            for c in ws[get_column_letter(i)][1:]:
                c.number_format = "#,##0"
        elif k == "precio_m2":
            for c in ws[get_column_letter(i)][1:]:
                c.number_format = "#,##0.00"
        elif k == "url":
            for c in ws[get_column_letter(i)][1:]:
                if c.value:
                    c.hyperlink = c.value
                    c.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    return ws


def datos_proyecto(p: Proyecto, cat: Catalogo, carpeta: Path, fichas: CacheFichas, estado: Estado):
    """Filas de listado deduplicadas (portal+id/url) y filas de hoteles del proyecto."""
    viviendas, hoteles, vistos = [], [], set()
    for clave in p.zonas:
        for u in cat.unidades(clave):
            for fila in filas_unidad(carpeta, u):
                k = (fila["portal"], fila.get("id") or fila["url"])
                if k in vistos:
                    continue
                vistos.add(k)
                f = _enriquecer(fila, u, cat, fichas)
                (hoteles if fila["portal"] == "booking" else viviendas).append(f)
    return viviendas, hoteles


def escribir_proyecto(p: Proyecto, cat: Catalogo, carpeta: Path, salida: Path, fichas: CacheFichas,
                      estado: Estado, comparte: dict[str, list[Proyecto]]) -> dict:
    viviendas, hoteles = datos_proyecto(p, cat, carpeta, fichas, estado)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    meta = [("Proyecto", p.nombre), ("Folios", ", ".join(p.folios)), ("País", p.pais),
            ("Responsable", p.responsable), ("Alcance (según hoja de proyectos)", p.alcance),
            ("Notas", p.notas), ("Anuncios de vivienda/terreno/comercial", len(viviendas)),
            ("Hoteles (Booking)", len(hoteles)),
            ("Fichas de detalle descargadas", sum(1 for f in viviendas if f["_ficha"]))]
    for k, v in meta:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.append([])
    ws.append(["Zona", "Fuente", "Operación", "Tipo", "Anuncios", "Estado", "Detalle", "Link",
               "Zona compartida con"])
    for c in ws[ws.max_row]:
        c.fill, c.font = _CAB, _FUENTE_CAB
    for clave in p.zonas:
        otros = ", ".join(sorted({x.nombre for x in comparte[clave] if x is not p})) or "-"
        for u in cat.unidades(clave):
            e = estado.datos.get(u.id, {})
            est = e.get("estado", "sin ejecutar")
            det = e.get("detalle", "")
            if u.verificar:
                det = ("URL por verificar. " + det).strip()
            ws.append([cat.zonas[clave].get("nombre", clave), u.etiqueta, u.operacion, u.tipo,
                       e.get("anuncios", 0), est, det, u.url, otros])
            fila = ws[ws.max_row]
            if est != "ok":
                for c in fila:
                    c.fill = _ERROR
            elif u.verificar or not e.get("anuncios"):
                for c in fila:
                    c.fill = _AVISO
    for col, ancho in zip("ABCDEFGHI", (34, 30, 12, 22, 40, 12, 50, 60, 40)):
        ws.column_dimensions[col].width = ancho
    ws.column_dimensions["B"].width = 40

    if viviendas or not hoteles:
        _hoja(wb, "Lista de precios", PRECIOS, viviendas)
    if hoteles:
        _hoja(wb, "Hoteles", HOTELES, hoteles)
    if viviendas:
        claves = []
        for f in viviendas:
            for k in f["_ficha"]:
                if k not in claves:
                    claves.append(k)
        cols = FICHA_BASE + [(k.replace("_", " ").capitalize(), f"ficha::{k}",
                              60 if "descripcion" in k else 24) for k in claves]
        for f in viviendas:
            for k, v in f["_ficha"].items():
                f[f"ficha::{k}"] = v
        _hoja(wb, "Ficha completa", cols, viviendas)
    salida.mkdir(parents=True, exist_ok=True)
    ruta = salida / p.archivo
    wb.save(ruta)
    return {"proyecto": p.nombre, "folios": ", ".join(p.folios), "pais": p.pais, "archivo": ruta.name,
            "viviendas": len(viviendas), "hoteles": len(hoteles),
            "fichas": sum(1 for f in viviendas if f["_ficha"])}


def escribir_excels(cat: Catalogo, proyectos: list[Proyecto], carpeta: Path, salida: Path) -> list[Path]:
    fichas, estado = CacheFichas(carpeta), Estado(carpeta)
    comparte = zonas_unicas(cat.proyectos)
    resumen = [escribir_proyecto(p, cat, carpeta, salida, fichas, estado, comparte) for p in proyectos]

    wb = Workbook()
    ws = wb.active
    ws.title = "Proyectos"
    ws.append(["Folios", "Proyecto", "País", "Prioridad", "Anuncios", "Hoteles", "Fichas", "Zonas", "Archivo"])
    por_nombre = {p.nombre: p for p in proyectos}
    for r in resumen:
        p = por_nombre[r["proyecto"]]
        ws.append([r["folios"], r["proyecto"], r["pais"], p.prioridad, r["viviendas"], r["hoteles"],
                   r["fichas"], ", ".join(p.zonas) or "(sin trabajo de portales)", r["archivo"]])
    for c in ws[1]:
        c.fill, c.font = _CAB, _FUENTE_CAB
    z = wb.create_sheet("Zonas únicas")
    z.append(["Zona", "Fuente", "Portal", "Anuncios", "Estado", "Detalle", "Usada por (se scrapea una vez)"])
    for c in z[1]:
        c.fill, c.font = _CAB, _FUENTE_CAB
    for clave, ps in zonas_unicas(proyectos).items():
        for u in cat.unidades(clave):
            e = estado.datos.get(u.id, {})
            z.append([cat.zonas[clave].get("nombre", clave), u.etiqueta, u.portal, e.get("anuncios", 0),
                      e.get("estado", "sin ejecutar"), e.get("detalle", ""),
                      "; ".join(x.nombre for x in ps)])
    for hoja, anchos in ((ws, (22, 36, 6, 9, 10, 9, 8, 50, 52)), (z, (34, 34, 14, 10, 12, 60, 70))):
        for i, a in enumerate(anchos, 1):
            hoja.column_dimensions[get_column_letter(i)].width = a
        hoja.freeze_panes = "A2"
    ruta_indice = salida / "00_indice_proyectos.xlsx"
    wb.save(ruta_indice)
    return [ruta_indice] + [salida / r["archivo"] for r in resumen]


def nombre_carpeta(p: Proyecto) -> str:
    """Nombre de carpeta válido en Windows: '<folios> - <proyecto>'."""
    import re
    base = f"{' + '.join(p.folios)} - {p.nombre}"
    base = re.sub(r'[<>:"/\\|?*]', "", base).strip(" .")
    return base[:120]


def entregar(cat: Catalogo, proyectos: list[Proyecto], carpeta: Path, salida: Path, destino: Path) -> list[tuple[str, int, int]]:
    """Copia el Excel de cada proyecto CON datos a ``destino/<folios> - <proyecto>/``.
    Los proyectos sin anuncios (ni hoteles) no generan carpeta."""
    import shutil
    fichas, estado = CacheFichas(carpeta), Estado(carpeta)
    res = []
    for p in proyectos:
        viv, hot = datos_proyecto(p, cat, carpeta, fichas, estado)
        if not (viv or hot):
            continue
        origen = salida / p.archivo
        if not origen.exists():
            continue
        carp = destino / nombre_carpeta(p)
        carp.mkdir(parents=True, exist_ok=True)
        shutil.copy2(origen, carp / origen.name)
        res.append((carp.name, len(viv), len(hot)))
    return res
